"""Fill the department's project-poster template (A0 portrait) in the style of the coordinator's example
(WIA poster): title, tool-logo strip, Abstract / Objectives / Key innovations / Datasets on the left, a
native architecture diagram on the right, Results + References + Team in a band at the bottom.
Sans-serif throughout. Every number is a measured result (results/).

    python scripts/poster_charts.py   # the four charts
    python scripts/build_poster.py    # -> docs/Review-III/QueryGenie_Poster.pptx
"""
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "docs", "Templates", "project-poster-tempalte.pptx")
OUT = os.path.join(ROOT, "docs", "Review-III", "QueryGenie_Poster.pptx")
ASSETS = os.path.join(ROOT, "docs", "Review-III", "poster_assets")
ICON = lambda name, colour="blue": os.path.join(ASSETS, "icons", f"{name}_{colour}.png")
BRANDI = lambda name: os.path.join(ASSETS, "icons", f"brand_{name}.png")
BRAND = os.path.join(ROOT, "desktop", "ui", "brand", "mark.png")

RED, NAVY, BLUE, YELLOW, GREEN = "C00000", "002060", "0070C0", "FFC000", "92D050"
INK, MUTED, LINE, CARD = "1F2937", "4B5563", "C9D3E0", "FFFFFF"
# diagram roles: reproduced CodeS / our enhancement / the application
BL_F, BL_L = "E8F1FB", "2F6FB5"
AM_F, AM_L = "FFF4DC", "C98A00"
GR_F, GR_L = "E8F5E9", "2E7D32"
FONT = "Arial"

prs = Presentation(TEMPLATE)
s = prs.slides[0]
P = {p.placeholder_format.idx: p for p in s.placeholders}
rgb = lambda h: RGBColor.from_string(h)


def place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def run(p, text, size, bold=False, colour=INK, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, FONT
    r.font.underline = False
    r.font.color.rgb = rgb(colour)
    return r


def text(x, y, w, h, paras, size=26, colour=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None, spacing=0.25, margin=0.08):
    tb = shape or s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if shape is not None:
        place(shape, x, y, w, h)
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * spacing)
        p.space_before = Pt(0)          # the master adds 33 pt above every paragraph
        runs = para if isinstance(para, list) else [para]
        for rn in runs:
            t, f = (rn, {}) if isinstance(rn, str) else rn
            run(p, t, f.get("size", size), f.get("bold", False), f.get("colour", colour), f.get("italic", False))
    return tb


def bar(x, y, w, title):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.8))
    b.fill.solid(); b.fill.fore_color.rgb = rgb(RED); b.line.fill.background(); b.shadow.inherit = False
    text(x, y, w, 0.8, [(title, {"bold": True, "colour": "FFFFFF", "size": 38})], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    return b


def header(idx, x, y, w, title):
    """One of the template's own section bars, moved, in the same sans style as bar()."""
    text(x, y, w, 0.8, [(title, {"bold": True, "colour": "FFFFFF", "size": 38})], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=P[idx])


def hide(idx):
    P[idx].text_frame.clear()
    place(P[idx], 0.7, 7.6, 0.05, 0.05)


def card(x, y, w, h, fill=CARD, line=LINE, radius=0.05, dashed=False, width=1.5):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = radius
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
    c.line.color.rgb = rgb(line); c.line.width = Pt(width); c.shadow.inherit = False
    if dashed:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def pic(path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def circle(x, y, d, colour, label):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = rgb(colour); c.line.fill.background(); c.shadow.inherit = False
    text(x, y, d, d, [(label, {"bold": True, "colour": "FFFFFF", "size": 28})], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c, margin=0)


def arrow(x1, y1, x2, y2, colour="5B6B80", label=None, lx=None, ly=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(colour); c.line.width = Pt(3.5)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    if label:
        text(lx if lx is not None else x2 + 0.1, ly if ly is not None else (y1 + y2) / 2 - 0.3, 3.2, 0.55,
             [(label, {"bold": True, "size": 20, "colour": colour})])


def node(x, y, w, h, title, lines, icon, fill, line, icon_colour="blue", title_colour=NAVY):
    card(x, y, w, h, fill=fill, line=line, radius=0.08, width=2.5)
    pic(ICON(icon, icon_colour), x + 0.25, y + (h - 1.0) / 2, w=1.0)
    paras = [(title, {"bold": True, "size": 27, "colour": title_colour})] + [(l, {"size": 24, "colour": INK}) for l in lines]
    text(x + 1.4, y + 0.08, w - 1.55, h - 0.16, paras, anchor=MSO_ANCHOR.MIDDLE, spacing=0.06)


# ------------------------------------------------------------------ title
text(6.6, 1.1, 21.4, 3.4, [("QueryGenie", {"size": 96, "bold": True, "colour": NAVY}),
                           ("A Self-Correcting, Confidence-Aware Natural Language to SQL Interface", {"size": 40, "colour": NAVY}),
                           ("Ask a database in plain English — and be told when the answer can't be trusted", {"size": 27, "colour": MUTED, "italic": True})],
     align=PP_ALIGN.CENTER, shape=P[153], spacing=0.04)
pic(BRAND, 29.0, 1.25, h=3.0)
for shp in s.shapes:                                    # the template's STUDENTS / INNOVATION / ... labels: same sans font
    if shp.has_text_frame and shp.text_frame.text.strip() in ("STUDENTS", "INNOVATION", "RESEARCH", "IMPACT"):
        for p_ in shp.text_frame.paragraphs:
            for r_ in p_.runs:
                r_.font.name = FONT; r_.font.bold = True

for shp in s.shapes:
    if shp.has_text_frame and shp.text_frame.text.startswith("Department of Computer Science"):
        for p_ in shp.text_frame.paragraphs:
            for r_ in p_.runs:
                r_.font.name = FONT; r_.font.bold = True

# ------------------------------------------------------------------ tool strip (as in the example: the actual tools, with logos)
TOOLS = [(BRANDI("python"), "Python 3.13"), (BRANDI("pytorch"), "PyTorch"), (BRANDI("huggingface"), "Hugging Face"),
         (BRANDI("sqlite"), "SQLite"), (BRANDI("postgresql"), "PostgreSQL · MySQL"), (ICON("monitor-smartphone", "navy"), "macOS · Windows app"),
         (ICON("badge-check", "green"), "45 tests"), (BRANDI("github"), "Open source · GitHub")]
n, gap = len(TOOLS), 0.2
cw = (31.7 - gap * (n - 1)) / n
for k, (img, label) in enumerate(TOOLS):
    cx = 0.7 + k * (cw + gap)
    card(cx, 5.9, cw, 1.25, radius=0.18)
    pic(img, cx + 0.15, 6.17, w=0.7)
    text(cx + 0.9, 5.9, cw - 0.95, 1.25, [(label, {"size": 20, "bold": True, "colour": NAVY})], anchor=MSO_ANCHOR.MIDDLE, margin=0.04)

# ------------------------------------------------------------------ left column
L, W = 0.7, 15.6
header(11, L, 7.5, W, "Abstract")
card(L, 8.4, W, 5.55, dashed=True)
text(L + 0.15, 8.5, W - 0.3, 5.4, [
    [("QueryGenie", {"bold": True}),
     (" lets anyone ask a database questions in plain English. Text-to-SQL models answer every question with the same "
      "confidence, so a wrong answer looks exactly like a right one. We reproduce CodeS (SIGMOD 2024), an open-source "
      "text-to-SQL model, to 77.8% execution accuracy on Spider (paper: 77.9%), and add confidence-aware abstention: "
      "the system runs every candidate query the model writes and checks whether they agree. When they don't, it asks "
      "the user instead of guessing. It answers 90.4% of questions and is right on 83.3% of them, with a third fewer "
      "wrong answers — a statistically significant gain. It runs fully locally in a desktop app.", {})]], size=31)
hide(10)

header(20, L, 14.25, W, "Objectives")
OBJ = ["Reproduce CodeS-1B on Spider within the ±5–10% tolerance",
       "Find which parts of the CodeS pipeline matter (ablation)",
       "Make the system say when it is unsure — tested statistically",
       "Measure accuracy on realistically worded questions",
       "Deliver a desktop app that runs fully locally"]
COL = [BLUE, "2E7D32", YELLOW, RED, "7030A0"]
for k, o in enumerate(OBJ):
    y = 15.25 + k * 1.05
    circle(L + 0.15, y, 0.8, COL[k], str(k + 1))
    text(L + 1.15, y - 0.07, W - 1.25, 0.95, [o], size=27, anchor=MSO_ANCHOR.MIDDLE)
hide(96)

bar(L, 20.65, W, "Key Innovations & Contributions")
INNOV = [("shield-check", "1. Confidence-aware abstention", "Run all candidate queries, count how many agree; ask “Did you mean…?” when they don't"),
         ("badge-check", "2. Faithful, audited reproduction", "77.8% vs 77.9% with the official evaluator, after auditing our own harness"),
         ("refresh-cw", "3. Execution-guided self-correction", "A failed query goes back to the model with the database error"),
         ("lock-keyhole", "4. Private by design", "Open weights, runs on the user's computer — no schema sent to a third party")]
for k, (ic, title, desc) in enumerate(INNOV):
    y = 21.65 + k * 1.72
    pic(ICON(ic, "blue"), L + 0.2, y + 0.18, w=0.95)
    text(L + 1.4, y, W - 1.5, 1.68, [(title, {"bold": True, "size": 27, "colour": NAVY}), (desc, {"size": 23, "colour": INK})], spacing=0.04)

bar(L, 28.65, W, "Datasets (Hugging Face — approved source)")
tbl = s.shapes.add_table(4, 3, Inches(L), Inches(29.55), Inches(W), Inches(2.6)).table
for j, wdt in enumerate([4.6, 4.4, 6.6]):
    tbl.columns[j].width = Inches(wdt)
ROWS = [("Benchmark", "Questions", "Used for"), ("Spider dev", "1,034 · 20 databases", "Reproduction, ablation, enhancement"),
        ("Spider-Realistic", "508", "Questions without column names"), ("BIRD Mini-Dev", "500 · 11 databases", "Real-world databases (future work)")]
for i, row in enumerate(ROWS):
    for j, v in enumerate(row):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid(); c.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("EEF3F9" if i % 2 else CARD))
        for p_ in c.text_frame.paragraphs:
            for r_ in p_.runs:
                r_.font.size = Pt(21); r_.font.name = FONT; r_.font.bold = i == 0
                r_.font.color.rgb = rgb("FFFFFF" if i == 0 else INK)

# ------------------------------------------------------------------ right column: the architecture, drawn natively
R = 16.8
header(25, R, 7.5, W, "Proposed Methodology / System Architecture")
hide(26)
card(R, 8.4, W, 23.75, dashed=True)
cx = R + W / 2
node(R + 0.4, 8.75, 7.0, 2.3, "USER QUESTION", ["“Which student scored the", "highest marks in DBMS?”"], "message-circle-question", BL_F, BL_L)
node(R + 8.2, 8.75, 7.0, 2.3, "YOUR DATABASE", ["SQLite · PostgreSQL · MySQL", "opened read-only"], "database", BL_F, BL_L)
arrow(R + 3.9, 11.05, cx - 1.5, 11.75); arrow(R + 11.7, 11.05, cx + 1.5, 11.75)
node(R + 1.5, 11.8, 12.6, 2.75, "DATABASE PROMPT BUILDER", ["Tables, columns, types, keys, 2 sample values each",
                                                           "Value retriever: stored values that match the question"], "file-text", BL_F, BL_L)
arrow(cx, 14.55, cx, 15.15)
node(R + 1.5, 15.2, 12.6, 2.3, "CODES-1B (FINE-TUNED ON SPIDER)", ["Beam search → 4 candidate SQL queries, each scored"], "brain-circuit", BL_F, BL_L)
arrow(cx - 2.5, 17.5, R + 3.9, 18.1)
node(R + 0.4, 18.15, 7.0, 2.6, "EXECUTE ALL 4", ["Read-only; keep the first", "valid query (paper's rule)"], "play", BL_F, BL_L)
node(R + 8.2, 18.15, 7.0, 2.6, "SELF-CORRECTION", ["Failed query + error →", "the model repairs it"], "refresh-cw", GR_F, GR_L, icon_colour="green")
arrow(R + 7.4, 19.45, R + 8.2, 19.45, colour=GR_L)
arrow(R + 3.9, 20.75, cx - 1.5, 21.35)
node(R + 1.5, 21.4, 12.6, 2.4, "EXECUTION AGREEMENT  ·  OUR ENHANCEMENT", ["How many different queries return the same result (0–4)"], "list-checks", AM_F, AM_L, icon_colour="orange")
arrow(cx, 23.8, cx, 24.4, colour=AM_L)
node(R + 1.5, 24.45, 12.6, 2.2, "CONFIDENT ENOUGH?", ["Agreement vs a threshold calibrated on Spider (strictness mode)"], "gauge", AM_F, AM_L, icon_colour="orange")
arrow(cx - 2.5, 26.65, R + 3.9, 27.35, colour=GR_L, label="yes", lx=R + 2.0, ly=26.7)
arrow(cx + 2.5, 26.65, R + 11.7, 27.35, colour=AM_L, label="no → ask", lx=R + 12.0, ly=26.7)
node(R + 0.4, 27.4, 7.0, 2.6, "ANSWER", ["Table + automatic chart", "SQL + “N of 4 agree”"], "chart-column", GR_F, GR_L, icon_colour="green")
node(R + 8.2, 27.4, 7.0, 2.6, "“DID YOU MEAN…?”", ["No guess: shows the different", "readings to choose from"], "circle-help", AM_F, AM_L, icon_colour="orange")
legend = [(BL_F, BL_L, "Reproduced from CodeS"), (AM_F, AM_L, "Our enhancement"), (GR_F, GR_L, "Application")]
for k, (f_, l_, t_) in enumerate(legend):
    lx = R + 0.6 + k * 5.1
    card(lx, 30.55, 0.55, 0.55, fill=f_, line=l_, radius=0.15, width=2.5)
    text(lx + 0.7, 30.45, 4.3, 0.75, [(t_, {"size": 21, "colour": INK})], anchor=MSO_ANCHOR.MIDDLE)
text(R + 0.4, 31.25, W - 0.8, 0.8, [("Desktop app (macOS · Windows) · runs locally on open weights", {"size": 21, "bold": True, "colour": NAVY})], align=PP_ALIGN.CENTER)

# ------------------------------------------------------------------ bottom band: results | references | team
Y = 32.6
RW, FW, TW = 19.0, 7.25, 5.05
header(154, L, Y, RW, "Experimental Results")
hide(155)
cwid, chh = 8.6, 8.6 * 4.8 / 7.6                    # four charts, 2 x 2, across the whole results width
off = (RW - (2 * cwid + 0.15)) / 2
for k, name in enumerate(["chart_reproduction", "chart_enhancement", "chart_wrong", "chart_calibration"]):
    pic(os.path.join(ASSETS, name + ".png"), L + off + (k % 2) * (cwid + 0.15), Y + 0.95 + (k // 2) * (chh + 0.1), w=cwid)
ry = Y + 0.95 + 2 * chh + 0.15
text(L, ry, RW, 45.75 - ry, [
    [("Significant: ", {"bold": True}), ("paired t-test p = 7.7×10⁻¹⁰, Wilcoxon p = 9.8×10⁻⁴, bootstrap 95% CI +4.5 to +6.7 points.", {})],
    [("Robust: ", {"bold": True}), ("counting each different query only once, accuracy at 90% coverage is still 83.5%.", {})]], size=22, spacing=0.08)

FX = L + RW + 0.2
header(27, FX, Y, FW, "References")
REFS = ["Li, H., et al. (2024). CodeS: Towards building open-source language models for text-to-SQL. arXiv:2402.16347.",
        "Yu, T., et al. (2018). Spider: A large-scale human-labeled dataset for complex and cross-domain semantic parsing and text-to-SQL task. arXiv:1809.08887.",
        "Deng, X., et al. (2020). Structure-grounded pretraining for text-to-SQL. arXiv:2010.12773.",
        "Li, J., et al. (2023). Can LLM already serve as a database interface? A big bench for large-scale database grounded text-to-SQLs. arXiv:2305.03111.",
        "Pourreza, M., & Rafiei, D. (2023). DIN-SQL: Decomposed in-context learning of text-to-SQL with self-correction. arXiv:2304.11015.",
        "Wang, X., et al. (2022). Self-consistency improves chain of thought reasoning in language models. arXiv:2203.11171.",
        "Geifman, Y., & El-Yaniv, R. (2017). Selective classification for deep neural networks. arXiv:1705.08500.",
        "Zhong, R., Yu, T., & Klein, D. (2020). Semantic evaluation for text-to-SQL with distilled test suites. arXiv:2010.02840."]
card(FX, Y + 0.95, FW, 45.7 - Y - 0.95, dashed=True)
text(FX + 0.1, Y + 1.05, FW - 0.2, 45.6 - Y - 1.05, [[(f"[{k + 1}] ", {"bold": True}), (r, {})] for k, r in enumerate(REFS)],
     size=23, spacing=0.35)
hide(28)

TX = FX + FW + 0.2
header(29, TX, Y, TW, "Team Details")
card(TX, Y + 0.95, TW, 5.6, dashed=True)
text(TX + 0.1, Y + 1.05, TW - 0.2, 5.4, [
    ("Supervisor", {"bold": True, "colour": NAVY, "size": 27}), ("M Parameswar", {}), ("Associate Professor", {"size": 22, "colour": MUTED}),
    ("Team 18", {"bold": True, "colour": NAVY, "size": 27}),
    ("2320090050", {"size": 21, "bold": True, "colour": MUTED}), ("N V V S S Jayakanth Kamisetti", {"size": 22}),
    ("2320090009", {"size": 21, "bold": True, "colour": MUTED}), ("Boddupally Narendar", {"size": 22}),
    ("2320090018", {"size": 21, "bold": True, "colour": MUTED}), ("M Praneeth Reddy", {"size": 22})], size=25, spacing=0.05)
hide(30)
bar(TX, Y + 6.8, TW, "Future Work")
text(TX, Y + 7.75, TW, 45.7 - Y - 7.75, [
    "• CodeS-3B and more, more diverse candidate queries",
    "• Real-world databases (BIRD)",
    "• A platform for any database (Spark, Snowflake, NoSQL) with access control"], size=24, spacing=0.3)

prs.save(OUT)
print("saved", OUT)
