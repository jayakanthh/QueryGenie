"""Fill the department's Review-III template in place (letterhead, theme and layouts kept).

    python scripts/build_review3_deck.py
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = f"{ROOT}/docs/Templates/Review-III-Project-Phase-I.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else f"{ROOT}/docs/Review-III/QueryGenie_Review-III.pptx"
FIG = f"{ROOT}/docs/report-latex/images/figures"
DATE = "10/3/26"
REALITY = None          # filled in when RL-001 / BD-001 finish: list of rows for the reality table

prs = Presentation(TEMPLATE)
S = list(prs.slides)


def dup(slide, after):
    """Copy a slide's shapes onto a new slide with the same layout, placed right after `after`."""
    new = prs.slides.add_slide(slide.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in slide.shapes:
        new.shapes._spTree.insert_element_before(copy.deepcopy(shp._element), "p:extLst")
    ids = prs.slides._sldIdLst
    el = ids[-1]
    ids.remove(el)
    pos = [s.slide_id for s in prs.slides].index(after.slide_id)
    ids.insert(pos + 1, el)
    return new


def new_after(layout, after):
    """A fresh slide with `layout`, placed right after `after`."""
    new = prs.slides.add_slide(layout)
    ids = prs.slides._sldIdLst
    el = ids[-1]
    ids.remove(el)
    pos = [s.slide_id for s in prs.slides].index(after.slide_id)
    ids.insert(pos + 1, el)
    return new


def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    return None


def bullets(shape, items, size=18, sub=15):
    """items: str (level 0) or (str, level)."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        bold = None
        if text.startswith("**") and "**" in text[2:]:
            head, rest = text[2:].split("**", 1)
            r = p.add_run(); r.text = head; r.font.bold = True; r.font.size = Pt(size if lvl == 0 else sub)
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size if lvl == 0 else sub)
        else:
            r = p.add_run(); r.text = text; r.font.size = Pt(size if lvl == 0 else sub)


def footer(slide):
    for p in slide.placeholders:
        if p.placeholder_format.idx == 10:
            p.text_frame.text = DATE


def table(slide, rows, left, top, width, col_w=None, size=12, head_size=None, row_h=0.32):
    n, m = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n, m, Inches(left), Inches(top), Inches(width), Inches(row_h * n))
    t = shp.table
    if col_w:
        for j, w in enumerate(col_w):
            t.columns[j].width = Inches(w)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = str(v)
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(head_size or size) if i == 0 else Pt(size)
                    run.font.bold = (i == 0) or (isinstance(v, str) and v.startswith("Enhanced, tested"))
            c.margin_left = c.margin_right = Inches(0.05)
            c.margin_top = c.margin_bottom = Inches(0.02)
    return t


def picture(slide, path, left, top, width=None, height=None):
    return slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width) if width else None,
                                    Inches(height) if height else None)


# ---------------------------------------------------------------- 1 title
s = S[0]
for shp in s.shapes:
    if not shp.has_text_frame:
        continue
    t = shp.text_frame.text
    if t.startswith("Project Title"):
        shp.text_frame.text = "QueryGenie: A Self-Correcting Natural Language to SQL Interface"
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(34)
    elif t.startswith("Presented by"):
        lines = ["Presented by — Team 18", "2320090050 : N V V S S Jayakanth Kamisetti", "2320090009 : Boddupally Narendar",
                 "2320090018 : M Praneeth Reddy"]
        paras = shp.text_frame.paragraphs
        for p, line in zip(paras, lines):
            for r in p.runs[1:]:
                r.text = ""
            if p.runs:
                p.runs[0].text = line
    elif t.startswith("Supervisor Name"):
        p0 = shp.text_frame.paragraphs[0]
        for r in p0.runs[1:]:
            r.text = ""
        p0.runs[0].text = "M Parameswar"
        paras = shp.text_frame.paragraphs
        if len(paras) > 1 and paras[1].runs:                 # "<Associate/Assistant Professor>"
            for r in paras[1].runs[1:]:
                r.text = ""
            paras[1].runs[0].text = "Associate Professor"

# ---------------------------------------------------------------- 2 outline (kept as the department wrote it)
footer(S[1])

# ---------------------------------------------------------------- 3 problem statement
bullets(ph(S[2], 1), [
    "**Problem: **people who don't know SQL can't get answers out of a database without help.",
    "**Text-to-SQL** translates a plain-English question into SQL, runs it and returns the result.",
    "**The risk: **these systems answer every question with the same confidence, so a wrong table looks exactly like a right one.",
    ("Even CodeS, a strong open model, is wrong on about 1 in 5 Spider questions (77.9% accuracy).", 1),
    "**Most strong systems use closed models** (GPT-4, …): the database schema goes to a third party.",
    "**Our goal: **reproduce CodeS, an open model that runs locally, and make it say when it is unsure instead of answering wrongly.",
], size=19, sub=16)
footer(S[2])

# ---------------------------------------------------------------- 4-7 related papers (16 papers, results from each paper's arXiv abstract)
PAPERS = [
    ("Seq2SQL (Zhong et al., 2017)", "Querying databases without knowing SQL", "WikiSQL (80,654 examples; introduced)", "Seq2seq with RL rewards from query execution", "Execution accuracy 35.9% → 59.4%"),
    ("Spider (Yu et al., 2018)", "Complex, cross-domain text-to-SQL", "Spider: 10,181 questions, 200 databases, 138 domains", "Human-annotated benchmark; unseen databases at test time", "Best model then: 12.4% exact match"),
    ("RAT-SQL (Wang et al., 2019)", "Generalising to unseen schemas", "Spider", "Relation-aware self-attention for schema encoding and linking", "57.2% exact match; 65.6% with BERT"),
    ("Test-suite accuracy (Zhong et al., 2020)", "Metrics misjudge equivalent queries", "Spider and 10 other datasets", "Distilled test suites of databases", "Spider metric: 2.5% false negatives on average"),
    ("StruG / Spider-Realistic (Deng et al., 2020)", "Questions rarely name the columns", "Spider-Realistic (introduced) + 8 datasets", "Structure-grounded pre-training", "Beats all baselines on the realistic sets"),
    ("BRIDGE (Lin et al., 2020)", "Linking questions to database contents", "Spider, WikiSQL", "Schema + matched cell values encoded with BERT", "Spider 71.1% dev, 67.5% test"),
    ("PICARD (Scholak et al., 2021)", "Language models produce invalid SQL", "Spider, CoSQL", "Constrained decoding by incremental parsing", "Fine-tuned T5 → state of the art"),
    ("Self-consistency (Wang et al., 2022)", "Unreliable single (greedy) answers", "GSM8K, SVAMP, AQuA, StrategyQA, ARC", "Sample several reasoning paths, keep the most consistent answer", "GSM8K +17.9%"),
    ("RESDSQL (Li et al., 2023)", "Schema linking and SQL skeleton parsed together", "Spider, Spider-DK, -Syn, -Realistic", "Ranking-enhanced encoder, skeleton-aware decoder", "Strong accuracy and robustness"),
    ("DIN-SQL (Pourreza & Rafiei, 2023)", "Prompting lags fine-tuned models", "Spider, BIRD", "Decomposed prompting with self-correction", "Spider test 85.3% EX; BIRD 55.9%"),
    ("BIRD (Li et al., 2023)", "Benchmarks ignore large, dirty databases", "BIRD: 12,751 pairs, 95 databases, 33.4 GB", "Benchmark with external knowledge and values", "ChatGPT 40.08% vs human 92.96%"),
    ("DAIL-SQL (Gao et al., 2023)", "No systematic study of LLM prompting", "Spider", "Prompt engineering comparison; DAIL-SQL prompt", "Spider leaderboard 86.6% EX"),
    ("MAC-SQL (Wang et al., 2023)", "Huge databases, multi-step questions", "BIRD, Spider", "Multi-agent: decomposer, selector, refiner", "BIRD test 59.59% EX (GPT-4)"),
    ("CodeS (Li et al., 2024) — reproduced", "Dependence on closed LLMs", "Spider, BIRD, robustness sets", "Open 1B–15B models, incremental pre-training on SQL", "CodeS-1B 77.9% EX (Spider dev)"),
    ("CHESS (Talaei et al., 2024)", "Large catalogs and schemas", "BIRD", "Four LLM agents: retrieve, prune, generate, test", "BIRD test 71.10%"),
    ("CHASE-SQL (Pourreza et al., 2024)", "Choosing among SQL candidates", "BIRD", "Multi-path generation + learned candidate selection", "BIRD test 73.0%"),
]
paper_slides = [S[3], S[4]]
paper_slides.append(dup(S[4], paper_slides[-1]))
paper_slides.append(dup(S[4], paper_slides[-1]))
for k, slide in enumerate(paper_slides):
    for shp in slide.shapes:
        if shp.has_table:
            t = shp.table
            chunk = PAPERS[4 * k: 4 * k + 4]
            for i in range(1, 5):
                row = chunk[i - 1] if i - 1 < len(chunk) else ("", "", "", "", "")
                for j in range(5):
                    c = t.cell(i, j)
                    c.text = row[j]
                    for p in c.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(11)
                t.rows[i].height = Inches(0.82)
            for j in range(5):
                for p in t.cell(0, j).text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(14)
        elif shp.has_text_frame and shp.text_frame.text.startswith("Identify at least"):
            shp.text_frame.text = f"Papers {4 * k + 1}–{min(4 * k + 4, len(PAPERS))} of {len(PAPERS)} (all from arXiv); results as stated in each paper's abstract."
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(13)
    if k > 0 and not any(sh.has_text_frame and sh.text_frame.text.startswith("Papers") for sh in slide.shapes):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(10.8), Inches(0.4))
        tb.text_frame.text = f"Papers {4 * k + 1}–{min(4 * k + 4, len(PAPERS))} of {len(PAPERS)} (all from arXiv); results as stated in each paper's abstract."
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.size = Pt(13)
    footer(slide)

# ---------------------------------------------------------------- research gap
bullets(ph(S[5], 1), [
    "**No confidence. **Every system reviewed answers every question; none says when it is unsure or reports an accuracy–coverage trade-off.",
    "**Closed models. **The strongest systems send the schema to GPT-4-class services: cost, privacy and reproducibility problems.",
    "**Benchmark-to-reality gap. **Results are reported on Spider, whose questions often name the columns and whose databases are small; how the same model does on realistic wording and real databases is rarely measured.",
    "**Our answer: **use the 4 candidate queries CodeS already generates — if they agree, answer; if they disagree, ask the user.",
], size=18)
footer(S[5])

# ---------------------------------------------------------------- dataset
s = S[6]
bullets(ph(s, 1), [
    "**All data from Hugging Face** (approved source); no Kaggle.",
    "**Spider dev** — xlangai/spider: 1,034 questions, 20 unseen databases (2–11 tables). Reproduction, ablation, enhancement.",
    ("39.5% need a JOIN, 26.8% GROUP BY, 8.0% nested queries; difficulty: 248 easy, 446 medium, 174 hard, 166 extra hard.", 1),
    "**Spider-Realistic** — 508 Spider questions reworded without column names (same databases and answers).",
    "**BIRD Mini-Dev** — 500 questions on 11 real-world databases (finance, medicine, sport…), up to 571 MB, dirty values.",
], size=17, sub=14)
footer(s)

# ---------------------------------------------------------------- architecture
s = S[7]
body = ph(s, 1)
body._element.getparent().remove(body._element)
picture(s, f"{FIG}/final_pipeline.png", 0.9, 1.75, width=7.6)
tb = s.shapes.add_textbox(Inches(8.7), Inches(1.8), Inches(4.2), Inches(4.9))
bullets(tb, [
    "**Reproduced CodeS (blue)**", ("Prompt: schema, keys, sample values, matched values", 1),
    ("4 candidate queries (beam search)", 1), ("Pick first query that runs", 1),
    "**Enhancement (amber)**", ("Run all 4, count how many agree", 1), ("Low agreement → ask the user", 1),
    "**Application (green)**", ("Desktop app, runs locally", 1), ("Table, chart, SQL, self-correction", 1),
], size=15, sub=13)
footer(s)

# ---------------------------------------------------------------- the application (new slide after the architecture)
app = new_after(S[7].slide_layout, S[7])
ph(app, 0).text_frame.text = "Proposed Methodology — the QueryGenie app"
b = ph(app, 1)
b._element.getparent().remove(b._element)
for k, (img, cap) in enumerate([("app_answer.png", "Confident answer: 4 of 4 queries agree → table, chart, SQL"),
                                ("app_not_sure.png", "Abstention: queries disagree → it asks which reading you meant")]):
    x = 0.8 + k * 6.05
    picture(app, f"{FIG}/{img}", x, 1.75, width=5.85)
    c = app.shapes.add_textbox(Inches(x), Inches(1.75 + 5.85 * 1600 / 2560 + 0.08), Inches(5.85), Inches(0.5))
    c.text_frame.word_wrap = True
    c.text_frame.text = cap
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
note = app.shapes.add_textbox(Inches(0.8), Inches(6.05), Inches(11.9), Inches(0.8))
note.text_frame.word_wrap = True
bullets(note, ["Desktop app (macOS, Windows) · runs locally on open weights · chats, schema map, charts, self-correction · "
               "SQLite, PostgreSQL, MySQL · optional Ollama / cloud models"], size=13)
footer(app)

# ---------------------------------------------------------------- results (two slides)
s = S[8]
ph(s, 0).text_frame.text = "Experimental Results"
body = ph(s, 1)
body._element.getparent().remove(body._element)
lbl = s.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(6), Inches(0.4))
lbl.text_frame.text = "Reproduction — Execution Accuracy, Spider dev (official evaluator)"
lbl.text_frame.paragraphs[0].runs[0].font.size = Pt(15); lbl.text_frame.paragraphs[0].runs[0].font.bold = True
table(s, [["", "EX"], ["CodeS-1B, paper", "77.9%"], ["Our reproduction (T2-006)", "77.8% (804/1,034)"],
          ["Difference", "−0.1 pts ✓ within ±5%"]], 0.8, 2.1, 5.6, col_w=[3.2, 2.4], size=13)
lbl2 = s.shapes.add_textbox(Inches(0.8), Inches(3.75), Inches(6), Inches(0.4))
lbl2.text_frame.text = "Ablation — each component removed"
lbl2.text_frame.paragraphs[0].runs[0].font.size = Pt(15); lbl2.text_frame.paragraphs[0].runs[0].font.bold = True
table(s, [["Removed", "EX", "Change"], ["Foreign keys", "76.9%", "−0.9"], ["Sample column values", "76.4%", "−1.4"],
          ["Candidate selection", "76.0%", "−1.8"], ["Value retriever", "75.9%", "−1.9"]], 0.8, 4.2, 5.6, col_w=[3.0, 1.3, 1.3], size=13)
lbl3 = s.shapes.add_textbox(Inches(6.9), Inches(1.65), Inches(6), Inches(0.4))
lbl3.text_frame.text = "Enhancement — same 1,034 questions"
lbl3.text_frame.paragraphs[0].runs[0].font.size = Pt(15); lbl3.text_frame.paragraphs[0].runs[0].font.bold = True
table(s, [["System", "Answered", "Correct when answering", "Wrong"],
          ["Original CodeS-1B", "100%", "77.8%", "230"],
          ["Enhanced, tested setting", "90.4%", "83.3%", "156"],
          ["3 of 4 must agree", "62.8%", "86.1%", "90"],
          ["All 4 must agree", "39.8%", "88.3%", "48"]], 6.9, 2.1, 5.9, col_w=[2.2, 1.0, 1.7, 1.0], size=12)
note = s.shapes.add_textbox(Inches(6.9), Inches(4.35), Inches(5.9), Inches(2.4))
note.text_frame.word_wrap = True
bullets(note, [
    "**Significant: **paired t-test p = 7.7×10⁻¹⁰, Wilcoxon p = 9.8×10⁻⁴, bootstrap 95% CI +4.5 to +6.7 pts.",
    "**Wrong answers −32% **(230 → 156) by declining 9.6% of questions.",
    "**Calibrated: **1 of 4 agree → 56.8% right; 4 of 4 → 88.3%.",
], size=14)
footer(s)

s2 = dup(S[8], S[8])
for shp in list(s2.shapes):
    if shp.shape_type in (19, 17) or (shp.has_text_frame and not shp.is_placeholder):
        shp._element.getparent().remove(shp._element)
ph(s2, 0).text_frame.text = "Experimental Results (continued)"
picture(s2, f"{FIG}/accuracy_coverage_T2-006.png", 0.8, 1.7, width=6.2)
tb = s2.shapes.add_textbox(Inches(7.2), Inches(1.7), Inches(5.7), Inches(0.5))
tb.text_frame.text = "Accuracy vs coverage: declining the least-agreed questions raises accuracy on the rest."
tb.text_frame.word_wrap = True
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
lab = s2.shapes.add_textbox(Inches(7.2), Inches(2.55), Inches(5.7), Inches(0.4))
lab.text_frame.text = "Self-correction (app feature), measured on Spider dev"
lab.text_frame.paragraphs[0].runs[0].font.size = Pt(15); lab.text_frame.paragraphs[0].runs[0].font.bold = True
table(s2, [["", "Without", "With repair"], ["EX", "77.76%", "77.85%"], ["Failed queries repaired", "—", "61 of 377 now run"]],
      7.2, 3.0, 5.7, col_w=[2.5, 1.4, 1.8], size=12)
lab2 = s2.shapes.add_textbox(Inches(7.2), Inches(4.3), Inches(5.7), Inches(0.4))
lab2.text_frame.text = "Benchmark-to-reality gap (CodeS-1B, unchanged)"
lab2.text_frame.paragraphs[0].runs[0].font.size = Pt(15); lab2.text_frame.paragraphs[0].runs[0].font.bold = True
rows = REALITY or [["Set", "EX (paper)", "Answered / correct"], ["Spider dev", "77.8% (77.9%)", "90.4% / 83.3%"],
                   ["Spider-Realistic", "69.5% (70.9%)", "86.2% / 76.7%"], ["BIRD Mini-Dev", "TBD", "TBD"]]
table(s2, rows, 7.2, 4.75, 5.7, col_w=[2.2, 1.2, 2.3], size=12)
footer(s2)

# ---------------------------------------------------------------- limitations
bullets(ph(S[9], 1), [
    "**Overall accuracy is unchanged: **the system is right more often when it answers, by declining ~10% of questions; plain EX stays 77.8%.",
    "**Agreement is not proof: **11.7% of answers where all 4 queries agree are still wrong.",
    "**One calibration: **the threshold was tuned on Spider with CodeS-1B; other models in the app are not measured.",
    "**Single questions: **CodeS was trained on one question at a time, so follow-ups need a chat model.",
    "**Reproduction scope: **inference only (pre-training needs 8× A800 GPUs); test-suite accuracy TBD.",
], size=17)
footer(S[9])

# ---------------------------------------------------------------- conclusion and future work
bullets(ph(S[10], 1), [
    "**Reproduced **CodeS-1B: 77.8% vs 77.9% (−0.1 pts) on a free Colab T4, after auditing our own evaluation.",
    "**Enhanced **with confidence-aware abstention: 83.3% correct when answering, 32% fewer wrong answers, statistically significant.",
    "**Delivered **QueryGenie, a desktop app that runs locally and asks when it isn't sure.",
    "**Future work — research:** test-suite accuracy; CodeS-3B; a learned reranker (a correct query is among the 4 for 83.1% of questions); calibration on real-world data.",
    "**Future work — from app to platform** (local, so no data leaves the company):",
    ("Any database: SQL Server, Snowflake, Apache Spark; MongoDB and other NoSQL through a SQL gateway", 1),
    ("Per-database adaptation: value index per database; optional one-time fine-tuning (CodeS data augmentation)", 1),
    ("Access control: users and roles per database / table / column, with an audit log", 1),
    ("Serving at scale: quantised larger models on laptops; a shared server for many users", 1),
], size=16, sub=14)
footer(S[10])

# ---------------------------------------------------------------- references (APA, 8 pt)
REFS = [
    "Deng, X., Awadallah, A. H., Meek, C., Polozov, O., Sun, H., & Richardson, M. (2020). Structure-grounded pretraining for text-to-SQL. arXiv:2010.12773.",
    "Gao, D., et al. (2023). Text-to-SQL empowered by large language models: A benchmark evaluation. arXiv:2308.15363.",
    "Geifman, Y., & El-Yaniv, R. (2017). Selective classification for deep neural networks. arXiv:1705.08500.",
    "Li, H., et al. (2023). RESDSQL: Decoupling schema linking and skeleton parsing for text-to-SQL. arXiv:2302.05965.",
    "Li, H., et al. (2024). CodeS: Towards building open-source language models for text-to-SQL. arXiv:2402.16347.",
    "Li, J., et al. (2023). Can LLM already serve as a database interface? A big bench for large-scale database grounded text-to-SQLs. arXiv:2305.03111.",
    "Lin, X. V., Socher, R., & Xiong, C. (2020). Bridging textual and tabular data for cross-domain text-to-SQL semantic parsing. arXiv:2012.12627.",
    "Pourreza, M., & Rafiei, D. (2023). DIN-SQL: Decomposed in-context learning of text-to-SQL with self-correction. arXiv:2304.11015.",
    "Pourreza, M., et al. (2024). CHASE-SQL: Multi-path reasoning and preference optimized candidate selection in text-to-SQL. arXiv:2410.01943.",
    "Scholak, T., Schucher, N., & Bahdanau, D. (2021). PICARD: Parsing incrementally for constrained auto-regressive decoding from language models. arXiv:2109.05093.",
    "Talaei, S., et al. (2024). CHESS: Contextual harnessing for efficient SQL synthesis. arXiv:2405.16755.",
    "Wang, B., Shin, R., Liu, X., Polozov, O., & Richardson, M. (2019). RAT-SQL: Relation-aware schema encoding and linking for text-to-SQL parsers. arXiv:1911.04942.",
    "Wang, B., et al. (2023). MAC-SQL: A multi-agent collaborative framework for text-to-SQL. arXiv:2312.11242.",
    "Wang, X., et al. (2022). Self-consistency improves chain of thought reasoning in language models. arXiv:2203.11171.",
    "Yu, T., et al. (2018). Spider: A large-scale human-labeled dataset for complex and cross-domain semantic parsing and text-to-SQL task. arXiv:1809.08887.",
    "Zhong, R., Yu, T., & Klein, D. (2020). Semantic evaluation for text-to-SQL with distilled test suites. arXiv:2010.02840.",
    "Zhong, V., Xiong, C., & Socher, R. (2017). Seq2SQL: Generating structured queries from natural language using reinforcement learning. arXiv:1709.00103.",
]
body = ph(S[11], 1)
tf = body.text_frame
tf.clear()
for k, ref in enumerate(REFS):
    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
    r = p.add_run(); r.text = ref; r.font.size = Pt(8)
footer(S[11])

# ---------------------------------------------------------------- GitHub URL
s = S[12]
t = ph(s, 0)
t.text_frame.text = "GitHub URL"
tb = s.shapes.add_textbox(Inches(0.9), Inches(4.2), Inches(11.5), Inches(0.8))
tb.text_frame.text = "https://github.com/jayakanthh/QueryGenie"
r = tb.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(28); r.hyperlink.address = "https://github.com/jayakanthh/QueryGenie"
footer(s)
footer(S[13])

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
